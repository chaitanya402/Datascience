#!/usr/bin/env python3
"""
Script to generate a Microsoft Word document about stars as celestial objects.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_stars_document():
    """Create a Word document about stars."""
    # Create a new Document
    doc = Document()
    
    # Title
    title = doc.add_heading('Stars: Celestial Objects of the Universe', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. Introduction
    doc.add_heading('1. Introduction: Overview of Stars as Celestial Phenomena', 1)
    
    intro_text = [
        "Stars are luminous celestial objects composed primarily of hydrogen and helium that generate "
        "energy through nuclear fusion in their cores. These brilliant bodies have captivated humanity "
        "for millennia, serving as navigational guides, subjects of scientific study, and sources of wonder.",
        "",
        "As the fundamental building blocks of galaxies, stars play a crucial role in the structure and "
        "evolution of the universe. They are responsible for creating and distributing most of the chemical "
        "elements that make up planets and life itself through nucleosynthesis.",
        "",
        "Our Sun is a typical star, and understanding stars helps us comprehend our place in the cosmos "
        "and the conditions necessary for life to exist."
    ]
    
    for paragraph in intro_text:
        if paragraph:
            doc.add_paragraph(paragraph)
        else:
            doc.add_paragraph()
    
    # 2. Types of Stars
    doc.add_heading('2. Types of Stars', 1)
    
    doc.add_paragraph(
        "Stars are classified based on various characteristics including mass, temperature, luminosity, "
        "and spectral type. Here are the main categories:"
    )
    
    # Main-Sequence Stars
    doc.add_heading('Main-Sequence Stars', 2)
    p = doc.add_paragraph()
    p.add_run('• ').bold = True
    p.add_run(
        "These stars, including our Sun, are in the stable phase of their lives, fusing hydrogen "
        "into helium in their cores. They represent about 90% of all stars in the universe."
    )
    
    # Giants
    doc.add_heading('Giants and Supergiants', 2)
    p = doc.add_paragraph()
    p.add_run('• Giants: ').bold = True
    p.add_run(
        "Stars that have exhausted hydrogen in their cores and expanded significantly. "
        "Examples include Aldebaran and Arcturus."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Supergiants: ').bold = True
    p.add_run(
        "Extremely massive and luminous stars like Betelgeuse and Rigel. These are among "
        "the largest and brightest stars known."
    )
    
    # White Dwarfs
    doc.add_heading('White Dwarfs', 2)
    p = doc.add_paragraph()
    p.add_run('• ').bold = True
    p.add_run(
        "The remnants of low to medium mass stars after they've shed their outer layers. "
        "These extremely dense objects are supported by electron degeneracy pressure. "
        "Sirius B is a famous example."
    )
    
    # Neutron Stars
    doc.add_heading('Neutron Stars', 2)
    p = doc.add_paragraph()
    p.add_run('• ').bold = True
    p.add_run(
        "The collapsed cores of massive stars following supernova explosions. These incredibly "
        "dense objects are composed primarily of neutrons and can spin rapidly as pulsars."
    )
    
    # Other Types
    doc.add_heading('Other Stellar Objects', 2)
    p = doc.add_paragraph()
    p.add_run('• Red Dwarfs: ').bold = True
    p.add_run("Small, cool, and long-lived stars that make up the majority of stars in the Milky Way.")
    
    p = doc.add_paragraph()
    p.add_run('• Brown Dwarfs: ').bold = True
    p.add_run("Failed stars that never achieved sustained nuclear fusion.")
    
    p = doc.add_paragraph()
    p.add_run('• Black Holes: ').bold = True
    p.add_run(
        "The final state of the most massive stars, where gravity is so strong that not "
        "even light can escape."
    )
    
    # 3. The Life Cycle of a Star
    doc.add_heading('3. The Life Cycle of a Star: Stellar Evolution', 1)
    
    doc.add_paragraph(
        "Stars undergo a fascinating transformation throughout their lives, with the specific path "
        "determined primarily by their initial mass:"
    )
    
    # Birth
    doc.add_heading('Birth: Stellar Nebulae', 2)
    p = doc.add_paragraph()
    p.add_run('• Formation: ').bold = True
    p.add_run(
        "Stars are born in vast clouds of gas and dust called nebulae. Gravity causes these "
        "clouds to collapse, forming protostars."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Ignition: ').bold = True
    p.add_run(
        "When the core temperature reaches about 10 million Kelvin, nuclear fusion begins, "
        "and a star is born."
    )
    
    # Main Sequence
    doc.add_heading('Main Sequence: The Stable Years', 2)
    p = doc.add_paragraph()
    p.add_run('• Duration: ').bold = True
    p.add_run(
        "Stars spend most of their lives (about 90%) in this stable phase, fusing hydrogen "
        "into helium."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Balance: ').bold = True
    p.add_run(
        "The outward pressure from nuclear fusion balances the inward pull of gravity, "
        "maintaining equilibrium."
    )
    
    # Evolution
    doc.add_heading('Evolution: Giants and Beyond', 2)
    p = doc.add_paragraph()
    p.add_run('• Low to Medium Mass Stars: ').bold = True
    p.add_run(
        "Become red giants, shed outer layers as planetary nebulae, and end as white dwarfs."
    )
    
    p = doc.add_paragraph()
    p.add_run('• High Mass Stars: ').bold = True
    p.add_run(
        "Become supergiants, explode as supernovae, and leave behind neutron stars or black holes."
    )
    
    # Death
    doc.add_heading('Death: The Final Stages', 2)
    p = doc.add_paragraph()
    p.add_run('• White Dwarf Cooling: ').bold = True
    p.add_run("Slowly radiates remaining heat over billions of years.")
    
    p = doc.add_paragraph()
    p.add_run('• Supernova Explosion: ').bold = True
    p.add_run(
        "Massive stars end in spectacular explosions, briefly outshining entire galaxies and "
        "seeding space with heavy elements."
    )
    
    # 4. The Importance of Stars
    doc.add_heading('4. The Importance of Stars', 1)
    
    # Role in the Universe
    doc.add_heading('Role in the Universe', 2)
    p = doc.add_paragraph()
    p.add_run('• Element Formation: ').bold = True
    p.add_run(
        "Stars create all elements heavier than hydrogen and helium through nuclear fusion "
        "and supernova explosions. We are literally made of stardust."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Galaxy Structure: ').bold = True
    p.add_run(
        "Stars are the visible components of galaxies and their gravitational interactions "
        "shape galactic structure."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Energy Source: ').bold = True
    p.add_run(
        "Stars are the primary energy sources in the universe, powering planetary systems "
        "and enabling life."
    )
    
    # Importance for Earth
    doc.add_heading('Importance for Earth', 2)
    p = doc.add_paragraph()
    p.add_run('• The Sun: ').bold = True
    p.add_run(
        "Our star provides the energy necessary for life on Earth through photosynthesis "
        "and drives weather patterns and climate."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Navigation: ').bold = True
    p.add_run(
        "For millennia, stars have guided travelers and explorers across land and sea."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Timekeeping: ').bold = True
    p.add_run(
        "Stellar positions have been used to mark seasons and create calendars throughout "
        "human history."
    )
    
    # Cultural Significance
    doc.add_heading('Cultural and Scientific Significance', 2)
    p = doc.add_paragraph()
    p.add_run('• Mythology and Religion: ').bold = True
    p.add_run(
        "Stars feature prominently in cultures worldwide, inspiring myths, legends, and "
        "religious beliefs."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Scientific Understanding: ').bold = True
    p.add_run(
        "Studying stars has led to breakthrough discoveries in physics, chemistry, and our "
        "understanding of the universe's age and evolution."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Search for Life: ').bold = True
    p.add_run(
        "Stars with planetary systems are targets in the search for extraterrestrial life."
    )
    
    # 5. Observing Stars
    doc.add_heading('5. Observing Stars', 1)
    
    # Telescope Observations
    doc.add_heading('Telescope Observations', 2)
    p = doc.add_paragraph()
    p.add_run('• Amateur Astronomy: ').bold = True
    p.add_run(
        "With even small telescopes or binoculars, amateur astronomers can observe thousands "
        "of stars, star clusters, and nebulae."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Professional Observatories: ').bold = True
    p.add_run(
        "Ground-based observatories like Mauna Kea and space telescopes like Hubble and "
        "James Webb provide unprecedented views of stellar phenomena."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Spectroscopy: ').bold = True
    p.add_run(
        "By analyzing starlight, astronomers can determine a star's composition, temperature, "
        "velocity, and other properties."
    )
    
    # Famous Stars and Constellations
    doc.add_heading('Famous Stars and Constellations', 2)
    p = doc.add_paragraph()
    p.add_run('• Polaris (North Star): ').bold = True
    p.add_run("The navigational star that marks Earth's northern celestial pole.")
    
    p = doc.add_paragraph()
    p.add_run('• Sirius: ').bold = True
    p.add_run("The brightest star in Earth's night sky, located in Canis Major.")
    
    p = doc.add_paragraph()
    p.add_run('• Betelgeuse: ').bold = True
    p.add_run(
        "A red supergiant in Orion that could explode as a supernova within the next "
        "100,000 years."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Orion Constellation: ').bold = True
    p.add_run("One of the most recognizable patterns, visible from both hemispheres.")
    
    p = doc.add_paragraph()
    p.add_run('• Pleiades (Seven Sisters): ').bold = True
    p.add_run("A beautiful open star cluster in Taurus, visible to the naked eye.")
    
    # Ongoing Research
    doc.add_heading('Ongoing Research and Discoveries', 2)
    p = doc.add_paragraph()
    p.add_run('• Exoplanet Discovery: ').bold = True
    p.add_run(
        "Modern telescopes have discovered thousands of planets orbiting other stars, "
        "revolutionizing our understanding of planetary systems."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Gravitational Waves: ').bold = True
    p.add_run(
        "Detected from colliding neutron stars and black holes, providing new insights "
        "into stellar death."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Dark Matter Studies: ').bold = True
    p.add_run(
        "Stellar motion helps astronomers map the distribution of dark matter in galaxies."
    )
    
    p = doc.add_paragraph()
    p.add_run('• Stellar Chemistry: ').bold = True
    p.add_run(
        "Advanced spectroscopy reveals the detailed chemical composition of stars across "
        "the universe."
    )
    
    # Conclusion
    doc.add_paragraph()
    doc.add_paragraph()
    conclusion = doc.add_paragraph(
        "Stars remain at the forefront of astronomical research and continue to reveal "
        "secrets about the universe's past, present, and future. From their role in creating "
        "the elements necessary for life to their potential for harboring other worlds, stars "
        "are truly the heart of cosmic exploration and understanding."
    )
    conclusion.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    return doc

def main():
    """Main function to generate and save the document."""
    import os
    
    print("Generating Word document about stars...")
    
    # Create the document
    doc = create_stars_document()
    
    # Save the document in the current directory
    output_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Stars.docx')
    doc.save(output_file)
    
    print(f"Document successfully created: {output_file}")

if __name__ == "__main__":
    main()
